import os
import re
import json
import sqlite3
import math
import uuid
import datetime
from pathlib import Path
from typing import List, Dict, Any, Optional, Tuple

class VectorEngine:
    """Per-database engine handling text extraction, chunking, vector indexing and similarity search."""

    def __init__(self, db_path: Path):
        self.db_path = Path(db_path).resolve()
        self.db_path.parent.mkdir(parents=True, exist_ok=True)
        self._init_db()

    def _init_db(self):
        """Initialize SQLite database tables for documents and chunks."""
        with sqlite3.connect(self.db_path) as conn:
            cursor = conn.cursor()
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS documents (
                    id TEXT PRIMARY KEY,
                    filename TEXT NOT NULL,
                    file_type TEXT NOT NULL,
                    file_size_bytes INTEGER NOT NULL,
                    uploaded_at TEXT NOT NULL,
                    chunk_count INTEGER NOT NULL,
                    folder TEXT DEFAULT 'General'
                )
            """)
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS chunks (
                    id TEXT PRIMARY KEY,
                    document_id TEXT NOT NULL,
                    filename TEXT NOT NULL,
                    chunk_index INTEGER NOT NULL,
                    content TEXT NOT NULL,
                    embedding TEXT NOT NULL,
                    FOREIGN KEY(document_id) REFERENCES documents(id) ON DELETE CASCADE
                )
            """)
            conn.commit()

    def get_stats(self) -> Dict[str, Any]:
        """Return total document count, chunk count, and file size in bytes."""
        file_size = self.db_path.stat().st_size if self.db_path.exists() else 0
        with sqlite3.connect(self.db_path) as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT COUNT(*) FROM documents")
            doc_count = cursor.fetchone()[0]
            cursor.execute("SELECT COUNT(*) FROM chunks")
            chunk_count = cursor.fetchone()[0]
        return {
            "document_count": doc_count,
            "chunk_count": chunk_count,
            "file_size_bytes": file_size
        }

    # --- Document Text Extraction ---

    def extract_text(self, file_path: Path, filename: str) -> str:
        """Extract plain text from supported document types (PDF, TXT, DOCX, MD, CSV, JSON, HTML)."""
        ext = filename.split(".")[-1].lower() if "." in filename else ""

        if ext == "pdf":
            return self._extract_pdf(file_path)
        elif ext == "docx":
            return self._extract_docx(file_path)
        elif ext in ["html", "htm"]:
            return self._extract_html(file_path)
        elif ext == "json":
            return self._extract_json(file_path)
        elif ext == "csv":
            return self._extract_csv(file_path)
        else:
            return self._read_raw_text(file_path)

    def _read_raw_text(self, file_path: Path) -> str:
        try:
            return file_path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            return file_path.read_text(encoding="latin1", errors="ignore")

    def _extract_pdf(self, file_path: Path) -> str:
        try:
            import fitz
            doc = fitz.open(str(file_path))
            text_pages = [page.get_text() for page in doc if page.get_text().strip()]
            doc.close()
            if text_pages:
                return "\n\n".join(text_pages)
        except Exception:
            pass

        try:
            import pypdf
            reader = pypdf.PdfReader(str(file_path))
            pages = [page.extract_text() for page in reader.pages if page.extract_text()]
            if pages:
                return "\n\n".join(pages)
        except Exception:
            pass

        return ""

    def _extract_docx(self, file_path: Path) -> str:
        try:
            import docx
            doc = docx.Document(str(file_path))
            return "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
        except Exception:
            return ""

    def _extract_html(self, file_path: Path) -> str:
        content = self._read_raw_text(file_path)
        clean = re.sub(r"<script.*?>.*?</script>", "", content, flags=re.DOTALL | re.IGNORECASE)
        clean = re.sub(r"<style.*?>.*?</style>", "", clean, flags=re.DOTALL | re.IGNORECASE)
        clean = re.sub(r"<[^>]+>", " ", clean)
        return re.sub(r"\s+", " ", clean).strip()

    def _extract_json(self, file_path: Path) -> str:
        content = self._read_raw_text(file_path)
        try:
            parsed = json.loads(content)
            return json.dumps(parsed, indent=2)
        except Exception:
            return content

    def _extract_csv(self, file_path: Path) -> str:
        return self._read_raw_text(file_path)

    # --- Sentence-Window Chunking ---

    def chunk_text(self, text: str, chunk_size: int = 400, overlap: int = 80) -> List[str]:
        """Split text into sentence-aware sliding window chunks."""
        text = text.strip()
        if not text:
            return []

        sentences = re.split(r"(?<=[.!?])\s+", text)
        chunks = []
        current_chunk = []
        current_len = 0

        for sentence in sentences:
            sentence = sentence.strip()
            if not sentence:
                continue

            if current_len + len(sentence) > chunk_size and current_chunk:
                chunk_str = " ".join(current_chunk)
                chunks.append(chunk_str)

                overlap_chunk = []
                overlap_len = 0
                for s in reversed(current_chunk):
                    if overlap_len + len(s) <= overlap:
                        overlap_chunk.insert(0, s)
                        overlap_len += len(s)
                    else:
                        break
                current_chunk = overlap_chunk
                current_len = overlap_len

            current_chunk.append(sentence)
            current_len += len(sentence)

        if current_chunk:
            chunks.append(" ".join(current_chunk))

        return chunks

    # --- Vector Embedding & Cosine Similarity ---

    def _tokenize(self, text: str) -> List[str]:
        return re.findall(r"\b\w{2,}\b", text.lower())

    def _compute_vector(self, text: str) -> Dict[str, float]:
        tokens = self._tokenize(text)
        if not tokens:
            return {}
        tf = {}
        for t in tokens:
            tf[t] = tf.get(t, 0.0) + 1.0
        total = math.sqrt(sum(v * v for v in tf.values()))
        return {k: v / total for k, v in tf.items()} if total > 0 else {}

    def _cosine_similarity(self, vec1: Dict[str, float], vec2: Dict[str, float]) -> float:
        if not vec1 or not vec2:
            return 0.0
        intersection = set(vec1.keys()) & set(vec2.keys())
        dot = sum(vec1[k] * vec2[k] for k in intersection)
        return dot

    # --- Document & Ingestion Operations ---

    def ingest_document(self, file_path: Path, filename: str, folder: str = "General") -> Dict[str, Any]:
        """Extract, chunk, compute embeddings and persist document & chunks to SQLite."""
        file_bytes = file_path.stat().st_size
        text_content = self.extract_text(file_path, filename)
        chunks = self.chunk_text(text_content)

        doc_id = str(uuid.uuid4())
        file_type = filename.split(".")[-1].lower() if "." in filename else "txt"
        uploaded_at = datetime.datetime.utcnow().isoformat()
        folder_clean = folder.strip() if folder and folder.strip() else "General"

        with sqlite3.connect(self.db_path) as conn:
            cursor = conn.cursor()
            cursor.execute(
                "INSERT INTO documents VALUES (?, ?, ?, ?, ?, ?, ?)",
                (doc_id, filename, file_type, file_bytes, uploaded_at, len(chunks), folder_clean)
            )

            for idx, chunk_text in enumerate(chunks):
                chunk_id = str(uuid.uuid4())
                vector = self._compute_vector(chunk_text)
                vector_json = json.dumps(vector)
                cursor.execute(
                    "INSERT INTO chunks VALUES (?, ?, ?, ?, ?, ?)",
                    (chunk_id, doc_id, filename, idx, chunk_text, vector_json)
                )

            conn.commit()

        return {
            "id": doc_id,
            "filename": filename,
            "file_type": file_type,
            "file_size_bytes": file_bytes,
            "uploaded_at": uploaded_at,
            "chunk_count": len(chunks),
            "folder": folder_clean,
        }

    def list_documents(self) -> List[Dict[str, Any]]:
        """List all documents in this database."""
        with sqlite3.connect(self.db_path) as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT id, filename, file_type, file_size_bytes, uploaded_at, chunk_count, folder FROM documents ORDER BY uploaded_at DESC")
            rows = cursor.fetchall()
            return [
                {
                    "id": r[0],
                    "filename": r[1],
                    "file_type": r[2],
                    "file_size_bytes": r[3],
                    "uploaded_at": r[4],
                    "chunk_count": r[5],
                    "folder": r[6] or "General",
                }
                for r in rows
            ]

    def get_document_chunks(self, doc_id: str) -> List[Dict[str, Any]]:
        """Get all text chunks for a specific document ordered by chunk_index."""
        with sqlite3.connect(self.db_path) as conn:
            cursor = conn.cursor()
            cursor.execute(
                "SELECT id, document_id, filename, chunk_index, content FROM chunks WHERE document_id = ? ORDER BY chunk_index ASC",
                (doc_id,)
            )
            rows = cursor.fetchall()
            return [
                {
                    "chunk_id": r[0],
                    "document_id": r[1],
                    "filename": r[2],
                    "chunk_index": r[3],
                    "content": r[4]
                }
                for r in rows
            ]

    def delete_document(self, doc_id: str) -> bool:
        """Delete a document and its vector chunks."""
        with sqlite3.connect(self.db_path) as conn:
            cursor = conn.cursor()
            cursor.execute("DELETE FROM chunks WHERE document_id = ?", (doc_id,))
            cursor.execute("DELETE FROM documents WHERE id = ?", (doc_id,))
            conn.commit()
            return cursor.rowcount > 0


    def search_chunks(self, query: str, top_k: int = 4, doc_ids: Optional[List[str]] = None, min_score: float = 0.0) -> List[Dict[str, Any]]:
        """Similarity search over chunks in this database."""
        query_vec = self._compute_vector(query)
        if not query_vec:
            return []

        with sqlite3.connect(self.db_path) as conn:
            cursor = conn.cursor()
            if doc_ids:
                placeholders = ",".join(["?"] * len(doc_ids))
                cursor.execute(
                    f"SELECT c.id, c.document_id, c.filename, c.chunk_index, c.content, c.embedding, d.folder FROM chunks c JOIN documents d ON c.document_id = d.id WHERE c.document_id IN ({placeholders})",
                    doc_ids
                )
            else:
                cursor.execute("SELECT c.id, c.document_id, c.filename, c.chunk_index, c.content, c.embedding, d.folder FROM chunks c JOIN documents d ON c.document_id = d.id")
            rows = cursor.fetchall()

        results = []
        for r in rows:
            chunk_id, doc_id, filename, chunk_idx, content, emb_json, doc_folder = r
            emb_vec = json.loads(emb_json)
            score = self._cosine_similarity(query_vec, emb_vec)
            if score >= min_score and score > 0.0:
                results.append({
                    "chunk_id": chunk_id,
                    "document_id": doc_id,
                    "filename": filename,
                    "text": content,
                    "score": round(score, 4),
                    "chunk_index": chunk_idx,
                    "folder": doc_folder or "General",
                })

        results.sort(key=lambda x: x["score"], reverse=True)
        return results[:top_k]
